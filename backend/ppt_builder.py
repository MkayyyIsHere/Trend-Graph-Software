# =============================================================================
# PPT BUILDER  (unchanged)
# =============================================================================

import os

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.enum.shapes import PP_PLACEHOLDER
    _PPTX_AVAILABLE = True
except ImportError:
    _PPTX_AVAILABLE = False

class PPTBuilder:
    IMAGE_EXTS = {'.png', '.jpg', '.jpeg'}

    class NoPicturePlaceholderError(Exception):
        pass

    def __init__(self, logger=None):
        self.logger = logger

    def _log(self, msg, level='INFO'):
        if self.logger:
            getattr(self.logger, level.lower())(msg)
        else:
            print(f"[{level}] {msg}")

    def collect_images(self, image_folder):
        entries = []
        for fname in os.listdir(image_folder):
            if os.path.splitext(fname)[1].lower() in self.IMAGE_EXTS:
                full = os.path.join(image_folder, fname)
                entries.append((os.path.getmtime(full), fname))
        entries.sort(key=lambda x: (x[0], x[1]))
        return [e[1] for e in entries]

    def find_best_layout(self, prs):
        if not _PPTX_AVAILABLE:
            return None
        title_types = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
        pic_types   = {PP_PLACEHOLDER.PICTURE}
        obj_types   = {PP_PLACEHOLDER.OBJECT}
        body_types  = {PP_PLACEHOLDER.BODY}
        candidates  = {1: None, 2: None, 3: None, 4: None}
        for layout in prs.slide_layouts:
            ph_types  = {ph.placeholder_format.type for ph in layout.placeholders
                         if ph.placeholder_format}
            has_title = bool(ph_types & title_types)
            if not has_title:
                continue
            if candidates[1] is None and ph_types & pic_types:  candidates[1] = layout
            if candidates[2] is None and ph_types & obj_types:   candidates[2] = layout
            if candidates[3] is None and ph_types & body_types:  candidates[3] = layout
            if candidates[4] is None:                            candidates[4] = layout
        chosen = candidates[1] or candidates[2] or candidates[3] or candidates[4]
        if chosen:
            self._log(f"Selected layout: '{chosen.name}'")
        return chosen

    def validate_layout_has_picture(self, layout) -> bool:
        if layout is None:
            return False
        title_types = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
        for ph in layout.placeholders:
            if ph.placeholder_format and ph.placeholder_format.type not in title_types:
                return True
        return False

    def _find_title_ph(self, slide):
        title_types = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
        for ph in slide.placeholders:
            if ph.placeholder_format and ph.placeholder_format.type in title_types:
                return ph
        return None

    def _find_picture_ph(self, slide):
        title_types = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
        pic_order   = {PP_PLACEHOLDER.PICTURE: 0, PP_PLACEHOLDER.OBJECT: 1, PP_PLACEHOLDER.BODY: 2}
        best_ph, best_score = None, 99
        for ph in slide.placeholders:
            pf = ph.placeholder_format
            if not pf or pf.type in title_types:
                continue
            score = pic_order.get(pf.type, 3)
            if score < best_score:
                best_score = score
                best_ph    = ph
        return best_ph

    @staticmethod
    def _image_px(image_path):
        try:
            from PIL import Image as PILImage
            with PILImage.open(image_path) as img:
                return img.size
        except Exception:
            return (1920, 1080)

    @staticmethod
    def _resize_slide_canvas(prs, width_px, height_px, dpi=96):
        emu_per_inch     = 914400
        prs.slide_width  = int(width_px  / dpi * emu_per_inch)
        prs.slide_height = int(height_px / dpi * emu_per_inch)

    def _insert_image_auto_resize(self, slide, prs, image_path, img_w, img_h, dpi=96):
        self._resize_slide_canvas(prs, img_w, img_h, dpi)
        pic_ph = self._find_picture_ph(slide)
        if pic_ph:
            try:
                pic_ph._element.getparent().remove(pic_ph._element)
            except Exception:
                pass
        slide.shapes.add_picture(image_path, 0, 0, prs.slide_width, prs.slide_height)

    def _insert_image_in_placeholder(self, slide, ph, image_path, img_w, img_h):
        ph_left, ph_top   = ph.left,  ph.top
        ph_width, ph_height = ph.width, ph.height
        if img_w > 0 and img_h > 0:
            scale = min(ph_width / img_w, ph_height / img_h)
            new_w, new_h = int(img_w * scale), int(img_h * scale)
        else:
            new_w, new_h = ph_width, ph_height
        left = ph_left + (ph_width  - new_w) // 2
        top  = ph_top  + (ph_height - new_h) // 2
        try:
            ph._element.getparent().remove(ph._element)
        except Exception:
            pass
        slide.shapes.add_picture(image_path, left, top, new_w, new_h)

    def build(self, image_folder, template_path, output_path,
              auto_resize_slide=True, slide_dpi=96,
              layout_index=None, progress_callback=None):
        if not _PPTX_AVAILABLE:
            self._log("python-pptx not installed.", 'ERROR')
            return None
        images = self.collect_images(image_folder)
        if not images:
            self._log(f"No PNG/JPG images found in: {image_folder}", 'ERROR')
            return None
        try:
            prs = Presentation(template_path)
        except Exception as e:
            self._log(f"Failed to load template: {e}", 'ERROR')
            return None

        layout = None
        if layout_index is not None:
            try:
                layout = prs.slide_layouts[layout_index]
            except IndexError:
                pass
        if layout is None:
            layout = self.find_best_layout(prs) or prs.slide_layouts[0]

        if not auto_resize_slide and not self.validate_layout_has_picture(layout):
            raise PPTBuilder.NoPicturePlaceholderError(
                f"Template layout '{layout.name}' has no picture placeholder.")

        for idx, filename in enumerate(images):
            image_path = os.path.join(image_folder, filename)
            if progress_callback:
                progress_callback(idx + 1, len(images), filename)
            try:
                img_w, img_h = self._image_px(image_path)
                slide        = prs.slides.add_slide(layout)
                title_ph     = self._find_title_ph(slide)
                if title_ph:
                    title_ph.text = os.path.splitext(filename)[0]
                if auto_resize_slide:
                    self._insert_image_auto_resize(slide, prs, image_path, img_w, img_h, slide_dpi)
                else:
                    pic_ph = self._find_picture_ph(slide)
                    if pic_ph:
                        self._insert_image_in_placeholder(slide, pic_ph, image_path, img_w, img_h)
                self._log(f"  [{idx+1}/{len(images)}] '{filename}'  ({img_w}×{img_h} px)")
            except Exception as e:
                self._log(f"  [{idx+1}] ERROR processing '{filename}': {e}", 'ERROR')

        try:
            prs.save(output_path)
            self._log(f"PPTX saved: {output_path}")
            return output_path
        except Exception as e:
            self._log(f"Failed to save PPTX: {e}", 'ERROR')
            return None
